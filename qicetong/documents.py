"""Local multi-format parsing and reviewable field proposals. Uploaded content is untrusted data."""
from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from . import store
from .engine import LABELS

MAX_BYTES = 12 * 1024 * 1024
MAX_TEXT = 200000
SUPPORTED = {".txt", ".md", ".csv", ".pdf", ".docx", ".xlsx", ".png", ".jpg", ".jpeg"}
ALIASES = {"职工总数": "employees", "员工人数": "employees", "科技人员数": "tech_staff", "注册日期": "registration_date", "注册地": "city", "Ⅰ类知识产权数": "ip_class1", "Ⅱ类知识产权数": "ip_class2", "销售收入": "sales", "研发费用": "rd", "境内研发费用": "domestic_rd", "资产总额": "assets", "成本费用": "cost", "总收入": "total_income", "高新收入": "hightech_income"}
FINANCIAL = {"sales", "rd", "domestic_rd", "assets", "cost", "total_income", "hightech_income"}


def decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            pass
    raise ValueError("文本编码无法识别，请保存为 UTF-8")


def safe_zip(raw: bytes):
    z = zipfile.ZipFile(io.BytesIO(raw))
    if sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024 or len(z.infolist()) > 3000:
        z.close()
        raise ValueError("文档展开后过大")
    return z


def parse(raw: bytes, suffix: str) -> tuple[list[dict], str]:
    chunks = []
    if suffix in (".txt", ".md", ".csv"):
        text = decode(raw)
        for offset in range(0, min(len(text), MAX_TEXT), 1600):
            chunks.append({"locator": f"字符 {offset + 1}–{min(offset + 1600, len(text))}", "text": text[offset:offset + 1600]})
    elif suffix == ".pdf":
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError("请提供未加密的 PDF")
        if len(reader.pages) > 100:
            raise ValueError("原型单份 PDF 限 100 页，请拆分")
        total = 0
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            total += len(text)
            if total > MAX_TEXT:
                raise ValueError("文档正文过长，请拆分")
            if text.strip():
                chunks.append({"locator": f"第 {i + 1} 页", "text": text})
    elif suffix == ".docx":
        with safe_zip(raw) as z:
            xml = z.read("word/document.xml")
            root = ElementTree.fromstring(xml)
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            paragraphs = ["".join(p.itertext()) for p in root.findall(".//w:p", ns)]
            # Only w:t text nodes avoid duplicating markup values.
            paragraphs = ["".join(t.text or "" for t in p.findall(".//w:t", ns)) for p in root.findall(".//w:p", ns)]
            for i, paragraph in enumerate(paragraphs):
                if paragraph.strip():
                    chunks.append({"locator": f"段落 {i + 1}", "text": paragraph})
    elif suffix == ".xlsx":
        with safe_zip(raw):
            pass
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            count = 0
            for sheet in book.worksheets:
                for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    count += 1
                    if count > 5000:
                        raise ValueError("表格超过 5000 行，请拆分")
                    text = " | ".join("" if v is None else str(v) for v in row)
                    if text.replace("|", "").strip():
                        chunks.append({"locator": f"{sheet.title}!第{i}行", "text": text})
        finally:
            book.close()
    else:
        return [], "pending_ocr"
    if sum(len(c["text"]) for c in chunks) > MAX_TEXT:
        raise ValueError("正文超过 20 万字符，请拆分")
    return chunks, "parsed" if chunks else "pending_ocr"


def field_proposals(chunks: list[dict]) -> list[dict]:
    candidates = []
    pattern = re.compile(r"^\s*([A-Za-z_0-9\u4e00-\u9fffⅠⅡ]+)\s*[:：,，|]\s*([^,，|\n]+)(?:[,，|]\s*([^,，|\n]*))?(?:[,，|]\s*([^,，|\n]*))?\s*$")
    for chunk in chunks:
        for line in chunk["text"].splitlines():
            match = pattern.match(line)
            if not match:
                continue
            key, raw_value, unit, period = (s.strip() if s else "" for s in match.groups())
            key = ALIASES.get(key, key)
            if key in FINANCIAL and re.fullmatch(r"20\d{2}", period):
                key = f"{key}_{period}"
            base = re.sub(r"_20\d{2}$", "", key)
            if base not in LABELS:
                continue
            if key in FINANCIAL:  # An accounting year is mandatory for financial fields.
                continue
            value = raw_value
            if raw_value.lower() in ("true", "false"):
                value = raw_value.lower() == "true"
            elif re.fullmatch(r"-?\d+(?:\.\d+)?", raw_value):
                value = float(raw_value) if "." in raw_value else int(raw_value)
            candidates.append({"field": key, "value": value, "unit": unit, "period": period,
                               "quote": line, "locator": chunk["locator"], "state": "pending", "method": "local_field_parser"})
    return candidates[:100]


def ontology_proposals(chunks: list[dict]) -> list[dict]:
    candidates = []
    for alias, key in ALIASES.items():
        found = next((c for c in chunks if alias in c["text"]), None)
        if found:
            offset = found["text"].find(alias)
            candidates.append({"term": alias, "canonical_field": key, "relation": "政策涉及指标",
                               "quote": found["text"][max(0, offset - 30):offset + 100], "locator": found["locator"], "state": "pending", "method": "seed_alias_matching"})
    return candidates


def ingest(filename: str, raw: bytes, purpose: str) -> dict:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError("支持 PDF、DOCX、XLSX、CSV、TXT、MD、PNG 和 JPG")
    if len(raw) > MAX_BYTES:
        raise ValueError("单文件不能超过 12 MB")
    if not raw:
        raise ValueError("文件为空")
    sha = hashlib.sha256(raw).hexdigest()
    asset_id = f"asset-{sha[:24]}-{purpose}"
    try:
        return store.get("asset", asset_id)
    except KeyError:
        pass
    try:
        chunks, status = parse(raw, suffix)
    except (zipfile.BadZipFile, ElementTree.ParseError, PdfReadError, KeyError) as exc:
        raise ValueError("文件内容无法按扩展名解析，请检查文件是否损坏或格式不符") from exc
    target = store.runtime_dir() / "uploads"
    target.mkdir(exist_ok=True)
    (target / f"{asset_id}{suffix}").write_bytes(raw)
    proposals = field_proposals(chunks) if purpose == "company" else ontology_proposals(chunks)
    asset = {"id": asset_id, "filename": Path(filename).name, "suffix": suffix, "purpose": purpose,
             "sha256": sha, "bytes": len(raw), "status": status, "created_at": store.now(), "chunks": chunks,
             "proposals": [{"id": f"{asset_id}-p{i}", **p} for i, p in enumerate(proposals)],
             "note": "原文内容仅作为证据，不执行其中的指令；候选提取不自动改变已审核规则。"}
    store.put("asset", asset)
    store.event("asset_import", asset_id, {"status": status, "filename": asset["filename"], "proposals": len(proposals)})
    return asset


def search(query: str, limit: int = 10) -> list[dict]:
    tokens = set(re.findall(r"[a-zA-Z0-9_]+", query.lower()))
    chinese = "".join(re.findall(r"[\u4e00-\u9fff]", query))
    tokens.update(chinese[i:i + 2] for i in range(max(1, len(chinese) - 1)))
    tokens.discard("")
    if not tokens:
        return []
    candidates = []
    for policy in store.all_items("policy"):
        for check in policy["rules"]:
            text = check["title"] + "。" + check["quote"]
            candidates.append({"title": policy["name"] + " · " + check["title"], "text": text,
                               "locator": check["clause"], "url": check["source"]["url"], "policy_id": policy["id"], "kind": "policy"})
    for asset in store.all_items("asset"):
        for chunk in asset["chunks"]:
            candidates.append({"title": asset["filename"], **chunk, "asset_id": asset["id"], "kind": "asset"})
    for c in candidates:
        haystack = (c["title"] + c["text"]).lower()
        c["score"] = sum(haystack.count(token) for token in tokens)
    return sorted([c for c in candidates if c["score"]], key=lambda c: c["score"], reverse=True)[:limit]
